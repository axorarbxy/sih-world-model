from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).with_name("architecture_document.docx")
doc = Document(); sec = doc.sections[0]; sec.top_margin = Inches(.55); sec.bottom_margin = Inches(.55); sec.left_margin = Inches(.68); sec.right_margin = Inches(.68)
styles = doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos'); styles['Normal'].font.size=Pt(9.5)
for name, size in [('Title',22),('Heading 1',13),('Heading 2',10.5)]:
    styles[name].font.name='Aptos Display' if name=='Title' else 'Aptos'; styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'),styles[name].font.name); styles[name].font.size=Pt(size); styles[name].font.color.rgb=RGBColor(0,0,0)
def p(text='', style=None, boldlead=None):
    para=doc.add_paragraph(style=style); para.paragraph_format.space_after=Pt(4); para.paragraph_format.line_spacing=1.04
    if boldlead and text.startswith(boldlead): para.add_run(boldlead).bold=True; para.add_run(text[len(boldlead):])
    else: para.add_run(text)
    return para
def shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)

title=p('Network Digital Twin Architecture', 'Title'); title.alignment=WD_ALIGN_PARAGRAPH.LEFT
p('Technical summary for SIH26153 AI based Network Attack Forecasting', 'Subtitle')
p('The prototype learns evolving network state from traffic telemetry, forecasts the next five state windows, and compares counterfactual mitigation futures. It runs locally on Windows and keeps raw traffic on the operator’s machine.')

p('System flow', 'Heading 1')
table=doc.add_table(rows=1, cols=5); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
for cell, label in zip(table.rows[0].cells, ['Traffic input','Feature state','World Model','Digital Twin','Defender view']):
    cell.text=label; shade(cell,'17365D'); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cell.paragraphs[0].runs: run.font.color.rgb=RGBColor(255,255,255); run.bold=True; run.font.size=Pt(8)
for label in ['PCAP or CIC and CTU CSV','Five second vectors','GRU Transformer attention','No action and mitigation','MITRE risk and evidence']:
    c=table.add_row().cells[len(table.rows[-1].cells)-len(table.rows[-1].cells)] if False else None
row=table.add_row().cells
for cell,label in zip(row,['Scapy and Pandas','Scaled NumPy matrix','Predicts next vector','Rolls out both futures','React and Plotly']): cell.text=label; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
for row in table.rows:
    for cell in row.cells:
        for para in cell.paragraphs: para.paragraph_format.space_after=Pt(1); para.alignment=WD_ALIGN_PARAGRAPH.CENTER

p('Data ingestion and state representation', 'Heading 1')
p('The API accepts PCAP, PCAPNG, and CSV flow records. Scapy reads packet captures while Pandas reads CIC IDS and CTU style CSV files. The extractor aggregates traffic into five second windows. Each window becomes a state vector with flow volume, TCP SYN ACK and RST activity, TTL statistics, TCP window size, payload statistics, inter arrival time, and SMB port 445 activity. Scikit learn StandardScaler normalizes the matrix before inference.')

p('Predictive model and stage mapping', 'Heading 1')
p('A GRU encodes short term sequence dynamics. A two layer temporal Transformer then models dependencies across the state sequence. Attention weights identify the past windows that influence the forecast. The model outputs the next state vector rather than an attack label. It rolls that output forward for K future windows. XGBoost maps each observed or predicted state to Reconnaissance, Initial Access, Lateral Movement, Command and Control, or Exfiltration.')

p('Counterfactual defense and explainability', 'Heading 1')
p('The Digital Twin produces Future A without intervention and Future B after a mitigation such as blocking SMB port 445. It calculates projected exfiltration risk for both trajectories and recommends the action only when it lowers the projected risk. The dashboard shows the probability timeline, ATT&CK heatmap, model attention, and ranked feature drivers for each selected forecast point. This gives the analyst evidence for the recommendation.')

p('Privacy, auditability, and deployment', 'Heading 1')
p('SQLite records prediction history and derived threat intelligence. The Adaptive Threat Genome writes a hash chained signature from selected flow features. It excludes packet payloads and IP addresses, allowing a future threat sharing mechanism without exposing raw traffic. FastAPI provides upload, history, and WebSocket streaming endpoints. React and Plotly provide the offline dashboard. Docker Compose supports reproducible deployment.')

p('Evaluation plan', 'Heading 1')
p('The repository includes a repeatable benchmark command. It reports precision, recall, F1 score, and false positive rate for logistic regression and a temporal context comparator. Synthetic results demonstrate the workflow only. Reportable results require labelled CIC IDS 2018 or CTU 13 state data and a documented train test split.')

doc.save(OUT)
print(OUT)
